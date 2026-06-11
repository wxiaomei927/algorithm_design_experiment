import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

# 创建工作簿
wb = Workbook()

# 创建表头字体
header_font = Font(bold=True)

# 1. 排序-两组100数据
ws = wb.create_sheet(title="排序-两组100数据")

# 读取两组数据
set1 = []
with open("data/sorting/random_set_1.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    next(reader)
    for row in reader:
        set1.append(int(row[1]))
        
set2 = []
with open("data/sorting/random_set_2.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    next(reader)
    for row in reader:
        set2.append(int(row[1]))

# 写入表头
ws.cell(row=1, column=1, value="index").font = header_font
ws.cell(row=1, column=2, value="random_set_1").font = header_font
ws.cell(row=1, column=3, value="random_set_2").font = header_font

# 写入数据
for i in range(100):
    ws.cell(row=i+2, column=1, value=i+1)
    ws.cell(row=i+2, column=2, value=set1[i])
    ws.cell(row=i+2, column=3, value=set2[i])

# 调整列宽
for col in ["A", "B", "C"]:
    ws.column_dimensions[col].width = 15

# 2. 排序-两组比较次数
ws = wb.create_sheet(title="排序-两组比较次数")
with open("results/sorting/two_sets_comparisons.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = header_font
for col in range(1, 6):
    ws.column_dimensions[get_column_letter(col)].width = 18

# 3. 排序-规模实验
ws = wb.create_sheet(title="排序-规模实验")
with open("results/sorting/scaled_sorting_results.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = header_font
for col in range(1, 10):
    ws.column_dimensions[get_column_letter(col)].width = 15

# 4. 排序-递归汇总
ws = wb.create_sheet(title="排序-递归汇总")
with open("results/sorting/recursion_summary.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = header_font
for col in range(1, 7):
    ws.column_dimensions[get_column_letter(col)].width = 15

# 5. 合并排序-递归明细
ws = wb.create_sheet(title="合并排序-递归明细")
with open("results/sorting/merge_recursion_sizes.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = header_font
for col in range(1, 6):
    ws.column_dimensions[get_column_letter(col)].width = 15

# 6. 快速排序-递归明细
ws = wb.create_sheet(title="快速排序-递归明细")
with open("results/sorting/quick_recursion_sizes.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = header_font
for col in range(1, 6):
    ws.column_dimensions[get_column_letter(col)].width = 15

# 7. 背包-正式1000物品
ws = wb.create_sheet(title="背包-正式1000物品")
with open("data/knapsack/required_items_1000.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = header_font
for col in range(1, 5):
    ws.column_dimensions[get_column_letter(col)].width = 15

# 8. 背包-正式规模结果
ws = wb.create_sheet(title="背包-正式规模结果")
with open("results/knapsack/required_scale_results.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = header_font
for col in range(1, 14):
    ws.column_dimensions[get_column_letter(col)].width = 15

# 9. 背包-解质量对比
ws = wb.create_sheet(title="背包-解质量对比")
with open("results/knapsack/required_solution_quality.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = header_font
for col in range(1, 6):
    ws.column_dimensions[get_column_letter(col)].width = 18

# 10. 背包-选择明细审计
ws = wb.create_sheet(title="背包-选择明细审计")
with open("results/knapsack/required_selected_items_audit.csv", "r", encoding="utf-8") as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 1:
                cell.font = header_font
for col in range(1, 12):
    ws.column_dimensions[get_column_letter(col)].width = 15

# 删除默认工作表
wb.remove(wb["Sheet"])

# 保存文件
wb.save("deliverables/学号-姓名-数据.xlsx")
print("Excel file generated successfully!")
print("Total sheets:", len(wb.sheetnames))
